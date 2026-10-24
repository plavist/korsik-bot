#!/usr/bin/env python3
"""Оценка качества поиска по базе знаний (RAG) на контрольных вопросах eval/rag_questions.jsonl.

Что измеряет (без вызова языковой модели — только эмбеддинги и поиск, поэтому быстро и дёшево):
  • вопросы, ответ на которые в базе есть: нашёлся ли нужный фрагмент (hit@1, hit@k, MRR) и ответит ли бот при пороге RAG_MIN_SCORE;
  • вопросы, ответа на которые нет: отсеет ли порог выдумку (бот промолчит и передаст вопрос HR).
Перебирает пороги и подсказывает лучший. Запуск на работающем стеке:
    python scripts/eval_rag.py --ingest http://127.0.0.1:8002 [--top-k 5] [--min-score 0.5] [--json result.json]
Проверка самого набора вопросов по локальным файлам без стека (текст берётся из PPTX/TXT/MD в папке):
    python scripts/eval_rag.py --validate base_docs
"""
import argparse
import json
import re
import sys
import urllib.request
from pathlib import Path

QUESTIONS = Path(__file__).resolve().parent.parent / "eval" / "rag_questions.jsonl"
THRESHOLDS = [round(0.3 + 0.05 * i, 2) for i in range(11)]  # 0.30 … 0.80


def load_questions(path: Path = QUESTIONS) -> list[dict]:
    rows = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    for r in rows:
        if not r.get("q") or not isinstance(r.get("answerable"), bool):
            raise ValueError(f"в вопросе нет q/answerable: {r}")
        if r["answerable"] and not (r.get("doc") and r.get("expect_any")):
            raise ValueError(f"у вопроса с ответом нужны doc и expect_any: {r['q']}")
    return rows


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").replace("\xa0", " ")).lower().strip()


def matches(hit: dict, q: dict) -> bool:
    """Фрагмент подходит, если он из нужного документа и содержит любую из ожидаемых фраз."""
    return norm(q["doc"]) in norm(hit["title"]) and any(norm(p) in norm(hit["text"]) for p in q["expect_any"])


def evaluate(results: list[tuple[dict, list[dict]]], threshold: float, top_k: int) -> dict:
    """results — [(вопрос, найденные фрагменты по убыванию score)]. Считает метрики при заданном пороге."""
    ans = [(q, h[:top_k]) for q, h in results if q["answerable"]]
    unans = [(q, h[:top_k]) for q, h in results if not q["answerable"]]
    ranks = [next((i + 1 for i, x in enumerate(h) if matches(x, q)), None) for q, h in ans]
    answered = sum(1 for (q, h) in ans if any(x["score"] >= threshold and matches(x, q) for x in h))
    refused = sum(1 for _, h in unans if not any(x["score"] >= threshold for x in h))
    return {
        "threshold": threshold, "top_k": top_k, "answerable": len(ans), "unanswerable": len(unans),
        "hit_at_1": sum(r == 1 for r in ranks) / len(ans) if ans else None,
        "hit_at_k": sum(r is not None for r in ranks) / len(ans) if ans else None,
        "mrr": sum(1 / r for r in ranks if r) / len(ans) if ans else None,
        "answered_correctly": answered / len(ans) if ans else None,          # бот ответит по нужному фрагменту
        "correct_refusals": refused / len(unans) if unans else None,         # бот промолчит там, где ответа нет
    }


def sweep(results, top_k: int) -> list[dict]:
    rows = [evaluate(results, t, top_k) for t in THRESHOLDS]
    for r in rows:
        r["balanced"] = ((r["answered_correctly"] or 0) + (r["correct_refusals"] or 0)) / 2
    return rows


def best(rows: list[dict]) -> dict:
    """Лучший порог по сбалансированной точности; при равенстве — тот, что чаще отвечает по базе."""
    return max(rows, key=lambda r: (round(r["balanced"], 6), r["answered_correctly"] or 0))


def search(ingest: str, query: str, top_k: int) -> list[dict]:
    req = urllib.request.Request(ingest.rstrip("/") + "/search", data=json.dumps({"query": query, "top_k": top_k}).encode(),
                                 headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=60) as r:
        return sorted(json.loads(r.read()), key=lambda x: -x["score"])


def pct(v) -> str:
    return "—" if v is None else f"{100 * v:.0f}%"


def report(rows: list[dict], current: dict, top: dict) -> str:
    out = ["порог  ответит по базе  промолчит там, где ответа нет  сбаланс.", "-" * 62]
    for r in rows:
        mark = "  ← сейчас" if r["threshold"] == current["threshold"] else ("  ← лучший" if r is top else "")
        out.append(f"{r['threshold']:.2f}   {pct(r['answered_correctly']):>14}   {pct(r['correct_refusals']):>29}   {pct(r['balanced']):>6}{mark}")
    out += ["", f"При пороге {current['threshold']:.2f}: hit@1 {pct(current['hit_at_1'])}, hit@{current['top_k']} {pct(current['hit_at_k'])}, "
                f"MRR {current['mrr']:.2f}; вопросов с ответом {current['answerable']}, без ответа {current['unanswerable']}.",
            f"Рекомендуемый RAG_MIN_SCORE по этому набору: {top['threshold']:.2f} (ответов по базе {pct(top['answered_correctly'])}, "
            f"верных молчаний {pct(top['correct_refusals'])})."]
    return "\n".join(out)


def validate(folder: Path, questions: list[dict]) -> list[str]:
    """Проверка набора: ожидаемые фразы действительно есть в текстах документов (PPTX/TXT/MD в папке). → список проблем."""
    from pptx import Presentation
    texts: dict[str, str] = {}
    for p in folder.iterdir():
        if p.suffix.lower() == ".pptx":
            def walk(shapes):
                for sh in shapes:
                    if sh.shape_type == 6:
                        yield from walk(sh.shapes)
                    elif sh.has_text_frame:
                        yield sh.text_frame.text
            texts[p.stem] = "\n".join(t for s in Presentation(str(p)).slides for t in walk(s.shapes))
        elif p.suffix.lower() in (".txt", ".md") and not p.name.upper().startswith("README"):
            texts[p.stem] = p.read_text(encoding="utf-8")
    problems = []
    for q in questions:
        if not q["answerable"]:
            continue
        doc = next((t for name, t in texts.items() if norm(q["doc"]) in norm(name)), None)
        if doc is None:
            problems.append(f"нет текста документа «{q['doc']}» (скан без текстового слоя? проверьте через стек): {q['q']}")
        elif not any(norm(p) in norm(doc) for p in q["expect_any"]):
            problems.append(f"в «{q['doc']}» нет ни одной ожидаемой фразы {q['expect_any']}: {q['q']}")
    return problems


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--ingest", default="http://127.0.0.1:8002", help="адрес сервиса ingest")
    ap.add_argument("--top-k", type=int, default=5)
    ap.add_argument("--min-score", type=float, default=0.5, help="текущий порог RAG_MIN_SCORE")
    ap.add_argument("--questions", default=str(QUESTIONS))
    ap.add_argument("--json", help="сохранить результат в файл")
    ap.add_argument("--validate", metavar="ПАПКА", help="проверить вопросы по локальным документам и выйти")
    args = ap.parse_args(argv)
    questions = load_questions(Path(args.questions))
    if args.validate:
        problems = validate(Path(args.validate), questions)
        print("\n".join(problems) if problems else f"Набор в порядке: {len(questions)} вопросов, ожидаемые фразы найдены в текстах.")
        return 1 if problems else 0
    results = [(q, search(args.ingest, q["q"], args.top_k)) for q in questions]
    rows = sweep(results, args.top_k)
    current = evaluate(results, args.min_score, args.top_k)
    top = best(rows)
    print(report(rows, current, top))
    if args.json:
        Path(args.json).write_text(json.dumps({"current": current, "sweep": rows, "best": top}, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0


if __name__ == "__main__":
    sys.exit(main())
