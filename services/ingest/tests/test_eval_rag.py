"""Скрипт оценки качества RAG (scripts/eval_rag.py): набор вопросов, метрики, перебор порогов, проверка набора по документам."""
import importlib.util
import json
from pathlib import Path

import pytest

SCRIPT = Path(__file__).resolve().parents[3] / "scripts" / "eval_rag.py"
spec = importlib.util.spec_from_file_location("eval_rag", SCRIPT)
ev = importlib.util.module_from_spec(spec)
spec.loader.exec_module(ev)

QA = {"q": "Сколько стоит подарок?", "doc": "Антикоррупционная политика", "expect_any": ["трех тысяч"], "answerable": True}
QB = {"q": "Какой пароль от wifi?", "answerable": False}


def hit(score, title="Антикоррупционная политика", text="не более трех тысяч рублей"):
    return {"score": score, "title": title, "text": text}


def test_real_question_set_is_well_formed():
    rows = ev.load_questions()
    assert sum(r["answerable"] for r in rows) >= 15 and sum(not r["answerable"] for r in rows) >= 5
    assert all(r["doc"] and r["expect_any"] for r in rows if r["answerable"])


def test_question_validation_errors(tmp_path):
    f = tmp_path / "q.jsonl"
    f.write_text(json.dumps({"q": "x", "answerable": True}) + "\n", encoding="utf-8")
    with pytest.raises(ValueError, match="doc и expect_any"):
        ev.load_questions(f)
    f.write_text(json.dumps({"answerable": False}) + "\n\n", encoding="utf-8")
    with pytest.raises(ValueError, match="q/answerable"):
        ev.load_questions(f)


def test_matching_ignores_case_spaces_and_requires_right_document():
    assert ev.matches(hit(0.9, text="Не более  ТРЕХ\xa0тысяч"), QA)
    assert not ev.matches(hit(0.9, title="Положение о коммерческой тайне"), QA)
    assert not ev.matches(hit(0.9, text="другой текст"), QA)


def test_metrics_thresholds_and_best_pick():
    results = [(QA, [hit(0.55, title="Положение о коммерческой тайне", text="x"), hit(0.52)]),   # нужный фрагмент вторым
               (QB, [hit(0.45, title="Антикоррупционная политика", text="случайный")])]          # выдумка с баллом 0.45
    low = ev.evaluate(results, 0.4, 5)
    assert low["hit_at_1"] == 0 and low["hit_at_k"] == 1 and low["mrr"] == 0.5
    assert low["answered_correctly"] == 1 and low["correct_refusals"] == 0                       # порог 0.4 пропускает выдумку
    high = ev.evaluate(results, 0.5, 5)
    assert high["answered_correctly"] == 1 and high["correct_refusals"] == 1                     # порог 0.5 — и ответит, и промолчит
    rows = ev.sweep(results, 5)
    assert len(rows) == len(ev.THRESHOLDS) and ev.best(rows)["threshold"] == 0.5               # лучший порог — по сбалансированной точности
    assert ev.evaluate([], 0.5, 5)["hit_at_1"] is None and ev.pct(None) == "—" and ev.pct(0.5) == "50%"


def test_report_marks_current_and_best_threshold():
    results = [(QA, [hit(0.7)]), (QB, [hit(0.2)])]
    rows = ev.sweep(results, 5)
    text = ev.report(rows, ev.evaluate(results, 0.5, 5), ev.best(rows))
    assert "← сейчас" in text and "Рекомендуемый RAG_MIN_SCORE" in text and "MRR 1.00" in text


def test_main_against_fake_search_and_json_output(tmp_path, monkeypatch, capsys):
    qs = tmp_path / "q.jsonl"
    qs.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in (QA, QB)), encoding="utf-8")
    monkeypatch.setattr(ev, "search", lambda ingest, q, k: [hit(0.8)] if q == QA["q"] else [hit(0.2)])
    out = tmp_path / "res.json"
    assert ev.main(["--questions", str(qs), "--min-score", "0.5", "--json", str(out)]) == 0
    assert "Рекомендуемый" in capsys.readouterr().out
    assert json.loads(out.read_text())["current"]["answered_correctly"] == 1


def test_search_sorts_hits_by_score(monkeypatch):
    class Resp:
        def __enter__(self): return self
        def __exit__(self, *a): pass
        def read(self): return json.dumps([hit(0.4), hit(0.9)]).encode()
    monkeypatch.setattr(ev.urllib.request, "urlopen", lambda req, timeout: Resp())
    assert [h["score"] for h in ev.search("http://x", "q", 5)] == [0.9, 0.4]


def test_validate_finds_missing_phrases_and_documents(tmp_path, capsys):
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Подарки не более трех тысяч рублей"
    prs.save(str(tmp_path / "Антикоррупционная политика.pptx"))
    (tmp_path / "README.md").write_text("не документ", encoding="utf-8")
    ok = {**QA}
    bad_phrase = {**QA, "q": "Другой вопрос", "expect_any": ["такой фразы нет"]}
    no_doc = {**QA, "q": "Про отпуск", "doc": "Правила внутреннего трудового распорядка"}
    assert ev.validate(tmp_path, [ok, QB]) == []
    problems = ev.validate(tmp_path, [ok, bad_phrase, no_doc, QB])
    assert len(problems) == 2 and "нет ни одной ожидаемой фразы" in problems[0] and "нет текста документа" in problems[1]
    (tmp_path / "Правила внутреннего трудового распорядка.txt").write_text("отпуск 28 календарных дней", encoding="utf-8")
    assert ev.validate(tmp_path, [{**no_doc, "expect_any": ["28 календарных"]}]) == []
    qs = tmp_path / "q.jsonl"
    qs.write_text(json.dumps(QA, ensure_ascii=False), encoding="utf-8")
    assert ev.main(["--questions", str(qs), "--validate", str(tmp_path)]) == 0 and "Набор в порядке" in capsys.readouterr().out
    qs.write_text(json.dumps(bad_phrase, ensure_ascii=False), encoding="utf-8")
    assert ev.main(["--questions", str(qs), "--validate", str(tmp_path)]) == 1


def test_document_update_rejects_unknown_columns():
    from app import db
    db.create("d1", "f.pdf", "Файл")
    with pytest.raises(ValueError, match="неизвестные поля"):
        db.update("d1", **{"status=status, title": "x"})                  # имя колонки не из белого списка → в SQL не попадает
    db.update("d1", status="ready", pages=3)
    assert db.get("d1")["status"] == "ready" and db.get("d1")["pages"] == 3
